from pathlib import Path
from typing import Optional
import json

try:
    import pypdf
except Exception:
    pypdf = None

try:
    import PyPDF2
except Exception:
    PyPDF2 = None

try:
    import docx
except Exception:
    docx = None

try:
    import pptx
except Exception:
    pptx = None

try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None


def load_document_text(path: str) -> dict:
    """
    Loads document and returns a dict with 'text' (full text) and 'pages' (list of page/chunk strings).
    Supports PDF, DOCX, PPTX, HTML, and plain text.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(path)

    suffix = p.suffix.lower()
    pages = []

    if suffix == ".pdf":
        pdf_module = pypdf or PyPDF2
        if pdf_module is None:
            raise RuntimeError("pypdf is required to load PDF files")
        with p.open("rb") as fh:
            reader = pdf_module.PdfReader(fh)
            for page in reader.pages:
                try:
                    pages.append(page.extract_text() or "")
                except Exception:
                    pages.append("")
        full = "\n".join(pages)
        if not full.strip():
            raise RuntimeError("PDF contains no extractable text")
        return {"text": full, "pages": pages}

    if suffix in (".docx",) and docx is not None:
        doc = docx.Document(str(p))
        paras = [para.text for para in doc.paragraphs if para.text]
        full = "\n".join(paras)
        # treat entire document as single page for now
        return {"text": full, "pages": [full]}

    if suffix in (".pptx",) and pptx is not None:
        prs = pptx.Presentation(str(p))
        for slide in prs.slides:
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
            pages.append("\n".join(texts))
        full = "\n".join(pages)
        return {"text": full, "pages": pages}

    if suffix in (".html", ".htm") and BeautifulSoup is not None:
        content = p.read_text(encoding="utf-8")
        soup = BeautifulSoup(content, "html.parser")
        # remove scripts/styles
        for s in soup(["script", "style"]):
            s.decompose()
        text = soup.get_text(separator="\n")
        return {"text": text, "pages": [text]}

    # fallback: read as plain text
    full = p.read_text(encoding="utf-8")
    return {"text": full, "pages": [full]}
